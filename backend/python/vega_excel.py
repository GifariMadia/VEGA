from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

FISCAL_MONTHS = ("Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar")
MONTH_NUMBERS = {name: index + 4 if index < 9 else index - 8 for index, name in enumerate(FISCAL_MONTHS)}
MONTH_PATTERN = re.compile(r"^(Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Jan|Feb|Mar)\s*['â€™]?\s*(\d{2,4})$", re.I)
MONEY_QUANTUM = Decimal("0.01")


class ExcelFileError(ValueError):
    pass


def _load_workbook(path: str | Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ExcelFileError("File tidak dapat dibuka sebagai workbook Excel yang valid.") from exc


def _decimal(value: Any, *, allow_blank: bool = False) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        if allow_blank:
            return Decimal("0")
        raise InvalidOperation
    if isinstance(value, bool):
        raise InvalidOperation
    try:
        amount = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise InvalidOperation from None
    if not amount.is_finite():
        raise InvalidOperation
    return amount


def _year_from_title(value: Any) -> int | None:
    if value is None:
        return None
    match = re.search(r"\bFY\s*'?\s*(20\d{2})\b", str(value), re.I)
    return int(match.group(1)) if match else None


def _month_header(value: Any) -> tuple[str, int] | None:
    if value is None:
        return None
    match = MONTH_PATTERN.match(str(value).strip())
    if not match:
        return None
    month = match.group(1).title()
    year = int(match.group(2))
    if year < 100:
        year += 2000
    return month, year


def _category_label(total_description: str) -> str:
    name = re.sub(r"^\s*total\s+", "", total_description, flags=re.I).strip()
    return " ".join(word if word.upper() in {"SGA", "COGS"} else word.title() for word in name.split())


def parse_budget_file(path: str | Path) -> dict[str, Any]:
    workbook = _load_workbook(path)
    if "MIS (FC)" not in workbook.sheetnames:
        workbook.close()
        return _refusal("BUDGET", "Sheet MIS (FC) tidak ditemukan.", "Gunakan sheet MIS (FC) sesuai template.")

    sheet = workbook["MIS (FC)"]
    row5 = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True), ())
    starts = [
        index
        for index, value in enumerate(row5)
        if str(value or "").strip().casefold() == "budget"
        and (index == 0 or str(row5[index - 1] or "").strip().casefold() != "budget")
    ]
    if len(starts) != 1:
        workbook.close()
        return _refusal("BUDGET", "Band Budget harus ada tepat satu kali.", "Satu blok Budget berisi 12 bulan pada baris 5.")

    month_start = starts[0]
    header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True), ())
    expected_months = list(FISCAL_MONTHS)
    detected_months = [_month_header(header[index] if index < len(header) else None) for index in range(month_start, month_start + 12)]
    detected_names = [item[0] for item in detected_months if item]
    if len(detected_months) != 12 or any(item is None for item in detected_months) or detected_names != expected_months:
        workbook.close()
        return _refusal("BUDGET", "Header bulan pada blok Budget tidak sesuai atau bergeser.", "Baris 6 harus berurutan Apr sampai Mar tepat di bawah band Budget.")

    anchor_headers = [str(header[index] or "").strip().casefold() if index < len(header) else "" for index in (0, 1, 6)]
    if anchor_headers != ["coa no.", "description", "fy'26 budget"]:
        workbook.close()
        return _refusal("BUDGET", "Header kolom COA, DESCRIPTION, atau FY'26 Budget tidak sesuai.", "Header harus berada pada kolom A, B, dan G.")

    fiscal_year = _year_from_title(sheet.cell(1, 1).value)
    if fiscal_year is None:
        fiscal_year = next((year for _, year in detected_months if year >= 2000), None)
    if fiscal_year is None:
        workbook.close()
        return _refusal("BUDGET", "Tahun fiskal tidak dapat dideteksi.", "Cantumkan tahun fiskal pada judul atau header bulan.")
    expected_headers = [(month, fiscal_year if period <= 9 else fiscal_year + 1) for period, month in enumerate(FISCAL_MONTHS, start=1)]
    if detected_months != expected_headers:
        workbook.close()
        return _refusal("BUDGET", "Tahun pada header bulan tidak sesuai dengan fiscal year.", "Apr-Des harus memakai tahun awal dan Jan-Mar tahun berikutnya.")

    rows_read = 0
    accepted_accounts: list[dict[str, Any]] = []
    budget_records: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen: set[str] = set()

    block_accounts: list[dict[str, Any]] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=7, values_only=True), start=7):
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue
        description = str(row[1] or "").strip() if len(row) > 1 else ""
        if "total" in description.casefold():
            category = _category_label(description)
            for pending in block_accounts:
                pending["category"] = category or None
            block_accounts = []
            continue
        raw_code = row[0] if row else None
        code = str(raw_code).strip() if raw_code is not None else ""
        if not re.fullmatch(r"\d{9}", code) or not description:
            continue
        rows_read += 1

        try:
            annual_amount = _decimal(row[6] if len(row) > 6 else None)
            monthly_amounts = [
                _decimal(row[month_start + offset] if month_start + offset < len(row) else None, allow_blank=True)
                for offset in range(12)
            ]
        except InvalidOperation:
            errors.append({"row": row_number, "issue": "Nilai budget tahunan atau bulanan bukan angka.", "expected": "Angka desimal yang valid."})
            continue

        monthly_total = sum(monthly_amounts, Decimal("0"))
        if abs(annual_amount - monthly_total) > Decimal("0.005"):
            errors.append({"row": row_number, "issue": f"Jumlah 12 bulan ({monthly_total}) tidak sama dengan FY'26 Budget ({annual_amount}).", "expected": "Selisih maksimum 0.005."})
            continue

        if code in seen:
            duplicates.append({"row": row_number, "coa_code": code, "issue": "Kode COA duplikat; kemunculan pertama dipakai."})
            continue
        seen.add(code)
        account = {"coa_code": code, "name": description, "category": None, "annual_amount": annual_amount, "source_row": row_number}
        accepted_accounts.append(account)
        block_accounts.append(account)
        for period, amount in enumerate(monthly_amounts, start=1):
            budget_records.append({"coa_code": code, "period": period, "amount": amount, "source_row": row_number})

    for pending in block_accounts:
        if "allocation" in pending["name"].casefold():
            pending["category"] = "Allocation"

    if not accepted_accounts:
        errors.append({"row": None, "issue": "Tidak ada baris COA budget yang dapat dimuat.", "expected": "Minimal satu kode COA 9 digit dengan deskripsi dan nilai valid."})

    result = {
        "success": not errors,
        "kind": "BUDGET",
        "fiscal_year": fiscal_year,
        "period": None,
        "sheet_read": "MIS (FC)",
        "rows_read": rows_read,
        "rows_accepted": len(accepted_accounts) if not errors else 0,
        "rows_rejected": len(errors),
        "total_amount": sum((account["annual_amount"] for account in accepted_accounts), Decimal("0")).quantize(MONEY_QUANTUM),
        "accounts": accepted_accounts,
        "records": budget_records if not errors else [],
        "duplicates": duplicates,
        "warnings": [],
        "errors": errors,
        "sample_rows": accepted_accounts[:5],
    }
    workbook.close()
    return result


def _refusal(kind: str, issue: str, expected: str) -> dict[str, Any]:
    return {
        "success": False,
        "kind": kind,
        "fiscal_year": None,
        "period": None,
        "sheet_read": None,
        "rows_read": 0,
        "rows_accepted": 0,
        "rows_rejected": 0,
        "total_amount": Decimal("0.00"),
        "accounts": [],
        "records": [],
        "duplicates": [],
        "warnings": [],
        "errors": [{"row": None, "issue": issue, "expected": expected}],
        "sample_rows": [],
    }


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return datetime.fromisoformat(value.strip()).date()
        except ValueError:
            return None
    return None


def parse_gl_file(path: str | Path, known_coa_codes: set[str] | None = None) -> dict[str, Any]:
    workbook = _load_workbook(path)
    if "CORE" not in workbook.sheetnames:
        workbook.close()
        return _refusal("GL", "Sheet CORE tidak ditemukan.", "Gunakan sheet CORE sesuai template.")
    sheet = workbook["CORE"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    expected_amount_headers = ["debits", "credits", "debits", "credits"]
    actual_amount_headers = [str(header[index] or "").strip().casefold() if index < len(header) else "" for index in range(11, 15)]
    if actual_amount_headers != expected_amount_headers:
        workbook.close()
        return _refusal("GL", "Header kolom L-O tidak sesuai atau bergeser.", "Kolom L-O harus Debits, Credits, Debits, Credits.")
    if len(header) < 19 or str(header[0] or "").strip() != "Pd." or str(header[2] or "").strip() != "Date" or str(header[3] or "").strip() != "Account Number":
        workbook.close()
        return _refusal("GL", "Header CORE tidak sesuai kontrak.", "Kolom A harus Pd., C Date, D Account Number; gunakan 19 kolom CORE.")

    rows_read = 0
    records: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    periods: set[int] = set()
    fiscal_years: set[int] = set()
    unknown: dict[str, list[int]] = {}

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue
        rows_read += 1
        account_number = str(row[3] or "").strip() if len(row) > 3 else ""
        if not account_number:
            skipped.append({"row": row_number, "issue": "Baris dilewati: Account Number kosong."})
            continue
        raw_period = row[0] if row else None
        if raw_period is None or not str(raw_period).strip():
            skipped.append({"row": row_number, "issue": "Baris dilewati: Pd. kosong."})
            continue
        try:
            period = int(str(raw_period).strip())
        except ValueError:
            errors.append({"row": row_number, "issue": f"Periode tidak valid: {raw_period!s}.", "expected": "Pd. berupa angka 01 sampai 12."})
            continue
        if period < 1 or period > 12:
            errors.append({"row": row_number, "issue": f"Periode di luar rentang: {period}.", "expected": "Pd. berupa angka 01 sampai 12."})
            continue

        try:
            native_debit = _decimal(row[11] if len(row) > 11 else None, allow_blank=True)
            native_credit = _decimal(row[12] if len(row) > 12 else None, allow_blank=True)
            converted_debit = _decimal(row[13] if len(row) > 13 else None)
            converted_credit = _decimal(row[14] if len(row) > 14 else None)
        except InvalidOperation:
            errors.append({"row": row_number, "issue": "Nilai debit/kredit bukan angka.", "expected": "Kolom L-O berisi angka; perbaiki nilai error Excel seperti #DIV/0!."})
            continue
        if converted_debit != 0 and converted_credit != 0:
            errors.append({"row": row_number, "issue": "Debit dan kredit USD sama-sama terisi pada satu baris.", "expected": "Hanya satu sisi debit atau kredit yang bernilai non-zero."})
            continue

        periods.add(period)
        account_parts = account_number.split("-")
        code = account_parts[0].strip()
        in_scope = len(account_parts) == 3 and account_parts[2].strip().upper() == "MIS000"
        section_missing = len(account_parts) == 2 and known_coa_codes is not None and code in known_coa_codes
        if not in_scope and not section_missing:
            continue

        txn_date = _date_value(row[2] if len(row) > 2 else None)
        if txn_date is None:
            warnings.append({"row": row_number, "issue": "Tanggal transaksi kosong atau tidak dikenali; fiscal year tidak dapat diverifikasi dari baris ini."})
            fiscal_year = None
        else:
            fiscal_year = txn_date.year if txn_date.month >= 4 else txn_date.year - 1
            fiscal_years.add(fiscal_year)
            expected_calendar_month = (period + 2) % 12 + 1
            if txn_date.month != expected_calendar_month:
                warnings.append({"row": row_number, "issue": f"Date {txn_date.isoformat()} tidak sesuai dengan Pd. {period:02d}; Pd. tetap menjadi sumber periode."})

        if len(account_parts) not in (2, 3) or not re.fullmatch(r"\d{9}", code):
            errors.append({"row": row_number, "issue": f"Account Number tidak valid: {account_number}.", "expected": "Kode COA 9 digit diikuti entity dan section opsional."})
            continue
        if known_coa_codes is not None and code not in known_coa_codes:
            unknown.setdefault(code, []).append(row_number)

        record = {
            "coa_code": code,
            "account_number": account_number,
            "section": account_parts[2].strip() if len(account_parts) == 3 else None,
            "description": str(row[4] or "").strip() if len(row) > 4 else "",
            "period": period,
            "fiscal_year": fiscal_year,
            "txn_date": txn_date,
            "amount": converted_debit - converted_credit,
            "currency": str(row[9] or "").strip().upper() if len(row) > 9 else "",
            "exchange_rate": _decimal(row[10], allow_blank=True) if len(row) > 10 else Decimal("0"),
            "debit_native": native_debit,
            "credit_native": native_credit,
            "reference": str(row[5] or "").strip() if len(row) > 5 else None,
            "vendor": str(row[6] or "").strip() if len(row) > 6 else None,
            "row_number": row_number,
        }
        records.append(record)

    if len(periods) > 1:
        errors.append({"row": None, "issue": f"File berisi lebih dari satu Pd.: {', '.join(f'{period:02d}' for period in sorted(periods))}.", "expected": "Satu file GL hanya boleh memuat satu periode."})
    if len(fiscal_years) > 1:
        errors.append({"row": None, "issue": "Tanggal GL menunjukkan lebih dari satu fiscal year.", "expected": "Satu file GL hanya boleh memuat satu fiscal year."})
    if records and not fiscal_years:
        errors.append({"row": None, "issue": "Fiscal year tidak dapat dideteksi dari tanggal transaksi.", "expected": "Tanggal transaksi yang valid diperlukan untuk menentukan tahun awal fiscal year."})
    for code, row_numbers in sorted(unknown.items()):
        errors.append({"row": row_numbers[0], "issue": f"Kode COA {code} tidak terdaftar.", "expected": "Daftarkan COA terlebih dahulu sebelum upload GL."})
    if not records and not errors:
        errors.append({"row": None, "issue": "Tidak ada baris GL untuk MIS000 yang dapat dimuat.", "expected": "Minimal satu transaksi in-scope pada sheet CORE."})

    accepted = len(records) if not errors else 0
    fiscal_year = next(iter(fiscal_years), None)
    period = next(iter(periods), None)
    total = sum((record["amount"] for record in records), Decimal("0")).quantize(MONEY_QUANTUM)
    result = {
        "success": not errors,
        "kind": "GL",
        "fiscal_year": fiscal_year,
        "period": period,
        "sheet_read": "CORE",
        "rows_read": rows_read,
        "rows_accepted": accepted,
        "rows_rejected": len(errors),
        "rows_skipped": skipped,
        "rows_filtered": rows_read - len(skipped) - len(records) - len(errors),
        "total_amount": total,
        "records": records if not errors else [],
        "unknown_coas": [{"coa_code": code, "rows": row_numbers} for code, row_numbers in sorted(unknown.items())],
        "warnings": warnings,
        "errors": errors,
        "sample_rows": records[:5],
    }
    workbook.close()
    return result