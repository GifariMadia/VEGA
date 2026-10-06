from __future__ import annotations

from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from validators import (
    FISCAL_MONTHS,
    build_summary,
    coerce_number,
    column_matches_alias,
    detect_reporting_period,
    detect_sheet_rows,
    find_first_value,
    normalize_header,
    read_excel_rows,
    resolve_month_name,
)


GL_REQUIRED_COLUMNS = [
    "posting date",
    "account number",
    "account description",
    "debit",
    "credit",
]


def _candidate_aliases() -> Dict[str, List[str]]:
    return {
        "account number": [
            "account number",
            "account no",
            "account code",
            "coa",
            "coa code",
            "gl account",
            "kode coa",
            "no coa",
            "account",
        ],
        "account description": [
            "account description",
            "account name",
            "description",
            "deskripsi",
            "nama akun",
            "item",
            "line item",
            "remarks",
        ],
        "posting date": [
            "posting date",
            "date",
            "tanggal",
            "document date",
        ],
        "entity": [
            "entity",
            "company",
            "business unit",
            "dept",
            "department",
            "nama entity",
        ],
        "debit": [
            "debit",
            "debits",
            "dr",
            "amount debit",
            "debit amount",
            "usd debit",
            "usd dr",
        ],
        "credit": [
            "credit",
            "credits",
            "cr",
            "amount credit",
            "credit amount",
            "usd credit",
            "usd cr",
        ],
        "doc number": [
            "document no",
            "doc no",
            "doc number",
            "document number",
            "reference",
            "ref",
            "voucher",
        ],
        "vendor": [
            "vendor",
            "vendor name",
            "supplier",
            "rekanan",
        ],
    }


def _extract_value(row: Dict[str, Any], aliases: List[str]) -> Any:
    for key, value in row.items():
        if column_matches_alias(key, aliases):
            return value
    return None


def _read_all_gl_rows(file_path: str) -> List[List[Any]]:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    all_rows: List[List[Any]] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            all_rows.append(list(row))
    return all_rows


def _index_matches(row_headers: List[str], aliases: List[str]) -> List[int]:
    matches: List[int] = []
    for idx, header in enumerate(row_headers):
        if not header:
            continue
        normalized_header = str(header).strip().lower().replace(" ", "")
        for alias in aliases:
            alias_value = str(alias).strip().lower().replace(" ", "")
            if len(alias_value) <= 2:
                continue
            if normalized_header == alias_value or normalized_header.endswith(alias_value) or normalized_header.startswith(alias_value):
                matches.append(idx)
                break
    return matches


def parse_gl_excel(file_path: str) -> Dict[str, Any]:
    rows = _read_all_gl_rows(file_path)
    if not rows:
        return {
            "file_type": "GL",
            "reporting_period": "FY2026/2027",
            "total_rows": 0,
            "accepted_rows": 0,
            "rejected_rows": 0,
            "missing_required_columns": [],
            "unmatched_coas": [],
            "validation_errors": ["Uploaded Excel file is empty."],
            "rows": [],
        }

    header_index, header_row = detect_sheet_rows(rows)
    if not header_row:
        return {
            "file_type": "GL",
            "reporting_period": "FY2026/2027",
            "total_rows": 0,
            "accepted_rows": 0,
            "rejected_rows": 0,
            "missing_required_columns": ["header row"],
            "unmatched_coas": [],
            "validation_errors": ["No readable header row detected."],
            "rows": [],
        }

    headers = [str(cell or "").strip() for cell in header_row]
    missing_required_columns = []
    for alias in ["posting date", "account number", "account description", "debit", "credit"]:
        if not any(column_matches_alias(header, _candidate_aliases()[alias]) for header in headers):
            missing_required_columns.append(alias)

    cleaned_rows: List[Dict[str, Any]] = []
    validation_errors: List[str] = []
    reporting_period = detect_reporting_period(rows)

    account_number_indexes = _index_matches(headers, _candidate_aliases()["account number"])
    account_description_indexes = _index_matches(headers, _candidate_aliases()["account description"])
    posting_date_indexes = _index_matches(headers, _candidate_aliases()["posting date"])
    debit_indexes = _index_matches(headers, _candidate_aliases()["debit"])
    credit_indexes = _index_matches(headers, _candidate_aliases()["credit"])
    doc_number_indexes = _index_matches(headers, _candidate_aliases()["doc number"])
    vendor_indexes = _index_matches(headers, _candidate_aliases()["vendor"])
    entity_indexes = _index_matches(headers, _candidate_aliases()["entity"])
    currency_indexes = _index_matches(headers, ["curr", "currency", "currency code"])

    for index in range(header_index + 1, len(rows)):
        row = rows[index]
        if not row or not any(cell is not None and str(cell).strip() not in {"", "-", "--"} for cell in row):
            continue

        if not account_number_indexes:
            validation_errors.append(f"Row {index + 1}: Missing COA / account number.")
            continue

        coa_code = row[account_number_indexes[0]] if account_number_indexes[0] < len(row) else None
        if coa_code is None or str(coa_code).strip() == "":
            validation_errors.append(f"Row {index + 1}: Missing COA / account number.")
            continue

        description = row[account_description_indexes[0]] if account_description_indexes and account_description_indexes[0] < len(row) else ""
        posting_date = row[posting_date_indexes[0]] if posting_date_indexes and posting_date_indexes[0] < len(row) else None

        debit_idx = debit_indexes[0] if debit_indexes else None
        credit_idx = credit_indexes[0] if credit_indexes else None
        usd_debit_idx = debit_indexes[1] if len(debit_indexes) > 1 else None
        usd_credit_idx = credit_indexes[1] if len(credit_indexes) > 1 else None

        local_debit = coerce_number(row[debit_idx]) if debit_idx is not None and debit_idx < len(row) else None
        local_credit = coerce_number(row[credit_idx]) if credit_idx is not None and credit_idx < len(row) else None
        usd_debit = coerce_number(row[usd_debit_idx]) if usd_debit_idx is not None and usd_debit_idx < len(row) else None
        usd_credit = coerce_number(row[usd_credit_idx]) if usd_credit_idx is not None and usd_credit_idx < len(row) else None

        if local_debit is None and usd_debit is not None:
            local_debit = usd_debit
        if local_credit is None and usd_credit is not None:
            local_credit = usd_credit

        if local_debit is None and local_credit is None and usd_debit is None and usd_credit is None:
            validation_errors.append(f"Row {index + 1}: Missing debit/credit values.")
            continue

        if local_debit is None:
            local_debit = 0.0
        if local_credit is None:
            local_credit = 0.0
        if usd_debit is None:
            usd_debit = local_debit
        if usd_credit is None:
            usd_credit = local_credit

        actual_amount = local_debit - local_credit
        period_month = None
        if posting_date is not None:
            try:
                month_name = str(posting_date).split("/")[-2] if "/" in str(posting_date) else None
                if month_name is not None:
                    period_month = resolve_month_name(month_name)
                else:
                    period_month = resolve_month_name(str(posting_date))
            except Exception:
                period_month = None

        if period_month is None:
            period_month = "Aug"

        currency = "USD"
        if currency_indexes and currency_indexes[0] < len(row):
            raw_currency = row[currency_indexes[0]]
            if raw_currency is not None and str(raw_currency).strip():
                currency = str(raw_currency).strip().upper()

        cleaned_rows.append(
            {
                "fiscal_year": reporting_period,
                "period_month": period_month,
                "coa_code": str(coa_code).strip(),
                "actual_amount": actual_amount,
                "debit": local_debit,
                "credit": local_credit,
                "usd_debit": usd_debit,
                "usd_credit": usd_credit,
                "doc_no": row[doc_number_indexes[0]] if doc_number_indexes and doc_number_indexes[0] < len(row) else None,
                "posting_date": posting_date,
                "vendor_name": row[vendor_indexes[0]] if vendor_indexes and vendor_indexes[0] < len(row) else None,
                "description": description,
                "department": row[entity_indexes[0]] if entity_indexes and entity_indexes[0] < len(row) else None,
                "section_code": None,
                "currency": currency,
                "upload_batch_id": None,
                "created_at": None,
                "updated_at": None,
            }
        )

    summary = build_summary(
        "GL",
        cleaned_rows,
        missing_required_columns,
        validation_errors,
        [],
        reporting_period,
    )
    summary["rows"] = cleaned_rows
    summary["sample_records"] = cleaned_rows[:5]
    return summary
