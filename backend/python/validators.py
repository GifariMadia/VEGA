from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

FISCAL_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
MONTH_ALIASES = {
    "jan": "Jan",
    "january": "Jan",
    "januari": "Jan",
    "feb": "Feb",
    "february": "Feb",
    "februari": "Feb",
    "mar": "Mar",
    "march": "Mar",
    "maret": "Mar",
    "apr": "Apr",
    "april": "Apr",
    "may": "May",
    "mei": "May",
    "jun": "Jun",
    "june": "Jun",
    "juni": "Jun",
    "jul": "Jul",
    "july": "Jul",
    "juli": "Jul",
    "aug": "Aug",
    "august": "Aug",
    "agus": "Aug",
    "agustus": "Aug",
    "sep": "Sep",
    "sept": "Sep",
    "september": "Sep",
    "oct": "Oct",
    "october": "Oct",
    "okt": "Oct",
    "oktober": "Oct",
    "nov": "Nov",
    "november": "Nov",
    "des": "Dec",
    "dec": "Dec",
    "december": "Dec",
    "desember": "Dec",
}


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def normalize_coa_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace('"', '').replace("'", '')
    text = text.replace(" ", "")
    return text.upper()


def coerce_number(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, bool):
            return None
        return float(value)

    text = str(value).strip()
    if text == "":
        return None
    if text.lower() in {"-", "--", "na", "n/a", "null", "none", "nan", "#na", "#n/a"}:
        return None

    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1].strip()
    text = text.replace("USD", "").replace("IDR", "").replace("Rp", "").replace("$", "")
    text = text.replace("\u00a0", " ").replace(",", "") if "." in text and "," in text else text
    text = text.replace(".", "") if "," in text and "." not in text else text
    text = text.replace(".", "") if text.count(".") > 1 else text
    text = text.replace(",", ".") if "," in text and "." not in text else text
    text = text.replace(" ", "")

    if text in {"", "-", "--"}:
        return None

    try:
        number = float(text)
        return -abs(number) if negative else number
    except ValueError:
        return None


def read_excel_rows(file_path: str) -> List[List[Any]]:
    file_path = str(file_path)
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    if not workbook.worksheets:
        raise ValueError(f"No worksheets found in {file_path}.")
    sheet = workbook[workbook.sheetnames[0]]
    return list(sheet.iter_rows(values_only=True))


def detect_sheet_rows(rows: Sequence[Sequence[Any]]) -> Tuple[int, List[str]]:
    best_index = 0
    best_score = -1
    best_row: List[str] = []

    for index, row in enumerate(rows[:60]):
        if not row:
            continue
        non_empty = [cell for cell in row if cell is not None and str(cell).strip() != ""]
        if len(non_empty) < 2:
            continue

        text = " ".join(str(cell or "") for cell in row).lower()
        score = 0
        if any(term in text for term in ["coa", "account", "description", "budget", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar"]):
            score += 3
        if "coa" in text or "account" in text:
            score += 2
        if "description" in text or "desc" in text:
            score += 2
        if any(term in text for term in ["apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec", "jan", "feb", "mar"]):
            score += 2
        if "budget" in text or "actual" in text:
            score += 1
        if score > best_score:
            best_score = score
            best_index = index
            best_row = [str(cell).strip() if cell is not None else "" for cell in row]

    if best_score < 0:
        return 0, []
    return best_index, best_row


def detect_reporting_period(rows: Sequence[Sequence[Any]]) -> str:
    for row in rows[:50]:
        joined = " ".join(str(cell or "") for cell in row)
        if re.search(r"fy\s*20\d{2}[/\\-]?20\d{2}", joined, re.I):
            return re.search(r"fy\s*20\d{2}[/\\-]?20\d{2}", joined, re.I).group(0).upper()
    for row in rows[:50]:
        text = " ".join(str(cell or "") for cell in row).lower()
        if "2026" in text and "2027" in text:
            return "FY2026/2027"
    return "FY2026/2027"


def resolve_month_name(cell: Any) -> Optional[str]:
    if cell is None:
        return None
    text = str(cell).strip().lower()
    if not text:
        return None
    if text in MONTH_ALIASES:
        return MONTH_ALIASES[text]
    for token, month in MONTH_ALIASES.items():
        if token in text:
            return month
    return None


def column_matches_alias(raw_header: Any, aliases: Iterable[str]) -> bool:
    normalized = normalize_header(raw_header)
    for alias in aliases:
        if normalize_header(alias) == normalized:
            return True
        if normalized and normalize_header(alias) in normalized:
            return True
    return False


def find_first_value(row: Dict[str, Any], aliases: Iterable[str]) -> Any:
    for key, value in row.items():
        if column_matches_alias(key, aliases):
            if value is not None and str(value).strip() != "":
                return value
    return None


def build_summary(
    file_type: str,
    rows: Sequence[Dict[str, Any]],
    missing_required_columns: List[str],
    validation_errors: List[str],
    unmatched_coas: List[str],
    reporting_period: str,
) -> Dict[str, Any]:
    accepted_rows = len(rows)
    rejected_rows = max(0, len(validation_errors))
    return {
        "file_type": file_type,
        "reporting_period": reporting_period,
        "total_rows": accepted_rows + rejected_rows,
        "accepted_rows": accepted_rows,
        "rejected_rows": rejected_rows,
        "missing_required_columns": missing_required_columns,
        "unmatched_coas": unmatched_coas,
        "validation_errors": validation_errors,
        "warnings": [],
    }
