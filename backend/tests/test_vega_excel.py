from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
import pytest

from backend.python.vega_excel import parse_budget_file, parse_gl_file

ROOT = Path(__file__).resolve().parents[2]
BUDGET_FILE = ROOT / "Docs" / "Source" / "Budget Dummy.xlsx"
GL_FILE = ROOT / "Docs" / "Source" / "GL Dummy.xlsx"


@pytest.fixture(scope="module")
def budget_result():
    return parse_budget_file(BUDGET_FILE)


@pytest.fixture(scope="module")
def budget_codes(budget_result):
    return {account["coa_code"] for account in budget_result["accounts"]}


@pytest.fixture(scope="module")
def gl_result(budget_codes):
    return parse_gl_file(GL_FILE, budget_codes | {"770102000", "770107001"})


def test_budget_fixture_obeys_locked_layout_and_totals(budget_result):
    result = budget_result

    assert result["success"] is True
    assert result["fiscal_year"] == 2026
    assert result["rows_read"] == 240
    assert result["rows_accepted"] == 238
    assert len(result["records"]) == 238 * 12
    assert str(result["total_amount"]) == "-67370.88"
    assert len(result["duplicates"]) == 2
    assert any(record["coa_code"] == "772502000" and record["amount"] < 0 for record in result["records"])


def test_gl_fixture_uses_core_pd_and_converted_columns(gl_result):
    result = gl_result
    assert result["success"] is True
    assert result["fiscal_year"] == 2026
    assert result["period"] == 3
    assert result["rows_accepted"] == 73
    assert len(result["records"]) == 73
    assert sum(record["section"] == "MIS000" for record in result["records"]) == 63
    assert sum(record["section"] is None for record in result["records"]) == 10
    assert str(result["total_amount"]) == "488981.16"
    assert result["unknown_coas"] == []


def test_gl_unknown_coas_are_refused(budget_codes):
    result = parse_gl_file(GL_FILE, budget_codes)

    assert result["success"] is False
    assert {item["coa_code"] for item in result["unknown_coas"]} == {"770102000", "770107001"}
    for code in ("770102000", "770107001"):
        assert any(code in error["issue"] and "tidak terdaftar" in error["issue"] for error in result["errors"])
    assert result["rows_accepted"] == 0
    assert result["records"] == []


def test_gl_shifted_converted_headers_are_refused(tmp_path):
    target = tmp_path / "shifted.xlsx"
    workbook = load_workbook(GL_FILE)
    workbook["CORE"]["N1"] = "Wrong Debit Header"
    workbook.save(target)

    result = parse_gl_file(target)

    assert result["success"] is False
    assert "bergeser" in result["errors"][0]["issue"]
    assert result["records"] == []


def test_budget_file_with_moved_band_is_refused(tmp_path):
    target = tmp_path / "shifted-budget.xlsx"
    workbook = load_workbook(BUDGET_FILE)
    workbook["MIS (FC)"].cell(6, 56).value = "May '26"
    workbook.save(target)

    result = parse_budget_file(target)

    assert result["success"] is False
    assert "bergeser" in result["errors"][0]["issue"]


@pytest.mark.parametrize("case", ["non_numeric", "both_sides", "mixed_periods", "negative", "date_crosscheck"])
def test_gl_edge_cases_small_workbooks(tmp_path, case):
    from datetime import datetime
    from openpyxl import Workbook
    from decimal import Decimal
    book=Workbook(); sheet=book.active; sheet.title="CORE"
    sheet.append(["Pd.","Srce.","Date","Account Number","Account Description","Reference","Vendor","Seq.","Batch-Entry","Curr.","Exch. Rate","Debits","Credits","Debits","Credits","Comment","FP Number","Doc. Number","Comment2"])
    row=["03","AP",datetime(2026,6,5),"999111222-A7744-MIS000","QA",None,None,1,None,"USD",1,12.5,0,12.5,0,None,None,None,None]
    if case == "non_numeric":row[13]="bad"
    if case == "both_sides":row[14]=1
    if case == "negative":row[13]=0;row[14]=12.5
    if case == "date_crosscheck":row[2]=datetime(2026,7,5)
    sheet.append(row)
    if case == "mixed_periods":
        second=list(row);second[0]="04";second[2]=datetime(2026,7,5);sheet.append(second)
    path=tmp_path/(case+".xlsx"); book.save(path)
    result=parse_gl_file(path,{"999111222"})
    if case in {"non_numeric","both_sides","mixed_periods"}:
        assert not result["success"] and not result["records"]
    else:
        assert result["success"]
        assert result["period"] == 3
        if case == "negative":assert result["total_amount"] == Decimal("-12.50")
        if case == "date_crosscheck":assert result["warnings"]
