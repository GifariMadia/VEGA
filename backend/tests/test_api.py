from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from time import monotonic

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.vega.dependencies import get_db
from backend.vega.main import create_app
from backend.vega.models import ActualEntry, AuditLog, Base, BudgetEntry, Coa, Department, UploadBatch, User
from backend.vega.routers import uploads
from backend.vega.security import hash_password


@pytest.fixture
def api(monkeypatch):
    from pathlib import Path

    monkeypatch.setenv("JWT_SECRET_KEY", "test-secret-that-is-long-enough-123456789")
    monkeypatch.setattr(uploads, "create_database_backup", lambda *_args, **_kwargs: Path("test-backup.sql"))
    import os
    test_url = os.getenv("VEGA_TEST_DATABASE_URL")
    if test_url:
        from sqlalchemy.engine import make_url
        if not (make_url(test_url).database or "").startswith("vega_test_"):
            raise RuntimeError("Integration tests require a dedicated vega_test_ database.")
        engine = create_engine(test_url)
        Base.metadata.drop_all(engine)
    else:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_db():
        session = factory()
        try:
            yield session
        finally:
            session.close()

    with factory.begin() as session:
        session.add(Department(code="MIS000", name="MIS Department"))
        session.add_all([
            User(username="admin", full_name="VEGA Admin", password_hash=hash_password("Admin-password-123!"), role="ADMIN", is_active=True),
            User(username="viewer", full_name="VEGA Viewer", password_hash=hash_password("Viewer-password-123!"), role="USER", is_active=True),
        ])

    app = create_app()
    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as client:
        def login(username: str, password: str) -> str:
            response = client.post("/api/v1/auth/login", json={"username": username, "password": password})
            assert response.status_code == 200, response.text
            return response.json()["data"]["access_token"]

        yield client, factory, login
    engine.dispose()


def _make_gl_file(path, account_number: str = "999999999-A7744-MIS000", debit: float = 12.5):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "CORE"
    headers = ["Pd.", "Srce.", "Date", "Account Number", "Account Description", "Reference", "Vendor", "Seq.", "Batch-Entry", "Curr.", "Exch. Rate", "Debits", "Credits", "Debits", "Credits", "Comment", "FP Number", "Doc. Number", "Comment2"]
    sheet.append(headers)
    row = ["03", "AP", datetime(2026, 6, 5), account_number, "Test account", "REF", "Vendor", 1, "BATCH", "USD", 1, debit, 0, debit, 0, None, None, None, None]
    sheet.append(row)
    workbook.save(path)
    return path


def _auth_headers(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def test_admin_and_user_permissions_and_coa_crud(api):
    client, factory, login = api
    admin_token = login("admin", "Admin-password-123!")
    user_token = login("viewer", "Viewer-password-123!")

    created = client.post("/api/v1/coa", headers=_auth_headers(admin_token), json={"code": "123456789", "name": "Test COA", "category": "Testing"})
    assert created.status_code == 201, created.text
    coa_id = created.json()["data"]["id"]
    assert client.get("/api/v1/coa", headers=_auth_headers(user_token)).status_code == 200
    denied_create = client.post("/api/v1/coa", headers=_auth_headers(user_token), json={"code": "123456788", "name": "Denied"})
    assert denied_create.status_code == 403

    updated = client.put(f"/api/v1/coa/{coa_id}", headers=_auth_headers(admin_token), json={"name": "Updated COA"})
    assert updated.status_code == 200
    deactivated = client.patch(f"/api/v1/coa/{coa_id}/status", headers=_auth_headers(admin_token), json={"is_active": False})
    assert deactivated.status_code == 200 and deactivated.json()["data"]["is_active"] is False
    deleted = client.delete(f"/api/v1/coa/{coa_id}", headers=_auth_headers(admin_token))
    assert deleted.status_code == 200

    upload_denied = client.post("/api/v1/uploads/preview", headers=_auth_headers(user_token), data={"kind": "GL"}, files={"file": ("tiny.xlsx", b"not a workbook", "application/octet-stream")})
    assert upload_denied.status_code == 403
    assert client.get("/api/v1/uploads/batches", headers=_auth_headers(user_token)).status_code == 200
    assert client.post(f"/api/v1/uploads/batches/1/cancel", headers=_auth_headers(user_token)).status_code == 403

    with factory() as db:
        assert db.get(Coa, coa_id).is_active is False


def test_upload_budget_gl_replace_unknown_refusal_and_cancel(api):
    from pathlib import Path

    client, factory, login = api
    admin_token = login("admin", "Admin-password-123!")
    user_token = login("viewer", "Viewer-password-123!")
    auth = _auth_headers(admin_token)
    root = Path(__file__).resolve().parents[2]

    budget_file = root / "Docs" / "Source" / "Budget Dummy.xlsx"
    budget_preview = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "BUDGET"}, files={"file": ("Budget Dummy.xlsx", budget_file.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert budget_preview.status_code == 200 and budget_preview.json()["success"] is True, budget_preview.text
    assert budget_preview.json()["data"]["rows_accepted"] == 238
    assert budget_preview.json()["data"]["total_amount"] == "-67370.88"
    budget_saved = client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": budget_preview.json()["data"]["preview_id"], "decision": "CONFIRM"})
    assert budget_saved.status_code == 200, budget_saved.text

    gl_file = root / "Docs" / "Source" / "GL Dummy.xlsx"
    gl_bytes = gl_file.read_bytes()
    unknown_preview = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("GL Dummy.xlsx", gl_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert unknown_preview.json()["success"] is False
    assert {"770102000", "770107001"}.issubset({error["issue"].split()[2] for error in unknown_preview.json()["errors"] if "tidak terdaftar" in error["issue"]})

    for code in ("770102000", "770107001"):
        response = client.post("/api/v1/coa", headers=auth, json={"code": code, "name": f"GL-derived {code}"})
        assert response.status_code == 201, response.text

    gl_preview = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("GL Dummy.xlsx", gl_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert gl_preview.json()["success"] is True, gl_preview.text
    assert gl_preview.json()["data"]["rows_accepted"] == 73
    assert gl_preview.json()["data"]["period"] == 3
    assert gl_preview.json()["data"]["total_amount"] == "488981.16"
    first_saved = client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": gl_preview.json()["data"]["preview_id"], "decision": "CONFIRM"})
    assert first_saved.status_code == 200, first_saved.text

    replacement_preview = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("GL Dummy.xlsx", gl_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert replacement_preview.json()["data"]["already_loaded"]["exists"] is True
    wrong_decision = client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": replacement_preview.json()["data"]["preview_id"], "decision": "CONFIRM"})
    assert wrong_decision.status_code == 409
    replaced = client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": replacement_preview.json()["data"]["preview_id"], "decision": "REPLACE"})
    assert replaced.status_code == 200, replaced.text

    with factory() as db:
        count_before_cancel = db.scalar(select(func.count()).select_from(ActualEntry))
        total_before_cancel = db.scalar(select(func.sum(ActualEntry.amount))).quantize(Decimal("0.01"))
        batches = db.scalars(select(UploadBatch).where(UploadBatch.kind == "GL").order_by(UploadBatch.id)).all()
        assert count_before_cancel == 73
        assert total_before_cancel == Decimal("488981.16")
        assert batches[0].status == "REPLACED" and batches[1].status == "ACTIVE"

    batch_id = replaced.json()["data"]["batch_id"]
    assert client.post(f"/api/v1/uploads/batches/{batch_id}/cancel", headers=_auth_headers(user_token)).status_code == 403
    cancelled = client.post(f"/api/v1/uploads/batches/{batch_id}/cancel", headers=auth)
    assert cancelled.status_code == 200, cancelled.text
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(ActualEntry)) == 0
        assert db.scalar(select(func.count()).select_from(BudgetEntry)) == 238 * 12
        assert db.get(UploadBatch, batch_id).status == "CANCELLED"


def test_corrupt_shifted_unknown_and_transaction_rollback(api, tmp_path):
    from backend.vega.routers.uploads import _previews

    client, factory, login = api
    admin_token = login("admin", "Admin-password-123!")
    auth = _auth_headers(admin_token)

    corrupt = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("broken.xlsx", b"not an xlsx", "application/octet-stream")})
    assert corrupt.status_code == 200 and corrupt.json()["success"] is False
    assert "workbook Excel yang valid" in corrupt.json()["errors"][0]["issue"]

    wrong_code = _make_gl_file(tmp_path / "unknown.xlsx")
    unknown = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("unknown.xlsx", wrong_code.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert unknown.json()["success"] is False
    assert "999999999" in unknown.json()["errors"][0]["issue"]

    from pathlib import Path

    root = Path(__file__).resolve().parents[2]
    shifted = load_workbook(root / "Docs" / "Source" / "GL Dummy.xlsx")
    shifted["CORE"]["N1"] = "Wrong Header"
    shifted_path = tmp_path / "shifted.xlsx"
    shifted.save(shifted_path)
    shifted_response = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("shifted.xlsx", shifted_path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert shifted_response.json()["success"] is False
    assert "bergeser" in shifted_response.json()["errors"][0]["issue"]

    with factory() as db:
        user = db.scalar(select(User).where(User.username == "admin"))
        coa = Coa(code="999999999", name="Test rollback COA", department_id=db.scalar(select(Department.id).where(Department.code == "MIS000")), is_active=True, is_gl_derived=False)
        db.add(coa)
        db.commit()
        user_id = user.id
        coa_id = coa.id

    preview_id = "rollback-test-preview-id-000000000000"
    _previews[preview_id] = {
        "created_at": monotonic(),
        "user_id": user_id,
        "filename": "rollback.xlsx",
        "parsed": {
            "kind": "GL", "fiscal_year": 2026, "period": 3, "rows_read": 1, "rows_accepted": 1,
            "rows_rejected": 0, "total_amount": Decimal("12.50"),
            "records": [{"coa_code": "999999999", "account_number": "999999999-A7744-MIS000", "section": "MIS000", "txn_date": None, "amount": Decimal("12.5"), "description": "fail", "currency": "USD", "exchange_rate": Decimal("1"), "debit_native": Decimal("12.5"), "credit_native": Decimal("0"), "reference": None, "vendor": None, "row_number": 2}],
        },
    }

    def fail_insert(_mapper, _connection, _target):
        raise SQLAlchemyError("simulated write failure")

    event.listen(ActualEntry, "before_insert", fail_insert)
    try:
        response = client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": preview_id, "decision": "CONFIRM"})
    finally:
        event.remove(ActualEntry, "before_insert", fail_insert)
        _previews.pop(preview_id, None)

    assert response.status_code == 500
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(ActualEntry)) == 0
        assert db.scalar(select(func.count()).select_from(UploadBatch)) == 0
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "UPLOAD")) == 0


def test_login_rate_limit(api, monkeypatch):
    from backend.vega.routers import auth as auth_router

    client, _, _ = api
    monkeypatch.setattr(auth_router, "LOGIN_ATTEMPT_LIMIT", 3)
    monkeypatch.setattr(auth_router, "LOGIN_ATTEMPT_WINDOW_SECONDS", 60)
    auth_router.FAILED_LOGINS.clear()

    for _ in range(3):
        response = client.post("/api/v1/auth/login", json={"username": "viewer", "password": "wrong-password"})
        assert response.status_code == 401, response.text

    locked = client.post("/api/v1/auth/login", json={"username": "viewer", "password": "wrong-password"})
    assert locked.status_code == 429, locked.text
    assert "terlalu banyak" in locked.json()["message"].lower()


def test_database_backup_helper_creates_dump_file(tmp_path, monkeypatch):
    import subprocess
    from pathlib import Path

    from backend.vega import backup as backup_module

    captured = {}

    def fake_run(cmd, capture_output, text, check, env):
        captured["cmd"] = cmd
        captured["env"] = env
        backup_path = Path(cmd[cmd.index("--file") + 1])
        backup_path.write_text("-- PostgreSQL database dump\n", encoding="utf-8")
        return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

    monkeypatch.setattr(backup_module.shutil, "which", lambda _name: "pg_dump")
    monkeypatch.setattr(backup_module.subprocess, "run", fake_run)
    result = backup_module.create_database_backup("postgresql://vega:secret@localhost:5432/vega", backup_dir=str(tmp_path), label="replace-batch-7")

    assert result.exists()
    assert result.name.startswith("vega-backup-")
    assert captured["cmd"][0].endswith("pg_dump")
    assert result.read_text(encoding="utf-8").startswith("-- PostgreSQL database dump")
    assert captured["env"]["PGPASSWORD"] == "secret"


def test_dashboard_aggregates_active_data_and_projects_year_end(api):
    client, factory, login = api
    token = login("viewer", "Viewer-password-123!")

    with factory.begin() as db:
        user_id = db.scalar(select(User.id).where(User.username == "admin"))
        department_id = db.scalar(select(Department.id).where(Department.code == "MIS000"))
        coa = Coa(code="123450001", name="Hardware", category="Hardware", department_id=department_id, is_active=True, is_gl_derived=False)
        db.add(coa)
        db.flush()
        budget_batch = UploadBatch(kind="BUDGET", filename="budget.xlsx", fy=2026, period=None, uploaded_by=user_id, rows_read=1, rows_imported=1, rows_rejected=0, total_amount=Decimal("100.00"), status="ACTIVE")
        gl_batch = UploadBatch(kind="GL", filename="gl.xlsx", fy=2026, period=1, uploaded_by=user_id, rows_read=1, rows_imported=1, rows_rejected=0, total_amount=Decimal("120.00"), status="ACTIVE")
        db.add_all([budget_batch, gl_batch])
        db.flush()
        db.add(BudgetEntry(fy=2026, period=1, coa_id=coa.id, amount=Decimal("100.00"), batch_id=budget_batch.id))
        db.add(ActualEntry(fy=2026, period=1, coa_id=coa.id, amount=Decimal("120.00"), account_number="123450001-MIS000", section="MIS000", txn_date=None, description="Test", currency="USD", exch_rate=Decimal("1"), debit_native=Decimal("120"), credit_native=Decimal("0"), row_no=2, batch_id=gl_batch.id))

    response = client.get("/api/v1/dashboard?fiscal_year=2026", headers=_auth_headers(token))
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["summary"]["budget"] == "100.00"
    assert data["summary"]["actual"] == "120.00"
    assert data["summary"]["variance"] == "-20.00"
    assert data["accounts"][0]["status"] == "OVER_BUDGET"
    assert data["monthly"][0] == {"period": 1, "month": "Apr", "budget": "100.00", "actual": "120.00"}
    assert data["projection"]["year_end_actual"] == "1440.00"


def test_user_management_and_two_role_rbac(api):
    client, factory, login = api
    admin = _auth_headers(login("admin", "Admin-password-123!"))
    viewer = _auth_headers(login("viewer", "Viewer-password-123!"))

    assert client.get("/api/v1/users", headers=viewer).status_code == 403
    assert client.post("/api/v1/users", headers=viewer, json={"username": "x-user", "full_name": "X", "role": "ADMIN", "password": "Long-enough-password-123!"}).status_code == 403
    assert client.get("/api/v1/users").status_code == 401

    created = client.post("/api/v1/users", headers=admin, json={"username": "new.viewer", "full_name": "New Viewer", "role": "USER", "password": "New-viewer-password-123!"})
    assert created.status_code == 201, created.text
    new_id = created.json()["data"]["id"]
    assert "password_hash" not in created.json()["data"]
    assert client.post("/api/v1/users", headers=admin, json={"username": "NEW.viewer", "full_name": "Dup", "role": "USER", "password": "New-viewer-password-123!"}).status_code == 409
    assert client.post("/api/v1/users", headers=admin, json={"username": "short-pw", "full_name": "S", "role": "USER", "password": "short"}).status_code == 422
    assert client.post("/api/v1/users", headers=admin, json={"username": "bad-role", "full_name": "S", "role": "ROOT", "password": "Long-enough-password-123!"}).status_code == 422

    new_login = client.post("/api/v1/auth/login", json={"username": "new.viewer", "password": "New-viewer-password-123!"})
    assert new_login.status_code == 200 and new_login.json()["data"]["user"]["role"] == "USER"
    new_headers = _auth_headers(new_login.json()["data"]["access_token"])
    assert client.get("/api/v1/dashboard?fiscal_year=2026", headers=new_headers).status_code == 200
    assert client.get("/api/v1/users", headers=new_headers).status_code == 403

    promoted = client.put(f"/api/v1/users/{new_id}", headers=admin, json={"role": "ADMIN"})
    assert promoted.status_code == 200 and promoted.json()["data"]["role"] == "ADMIN"
    assert client.get("/api/v1/users", headers=new_headers).status_code == 200

    deactivated = client.put(f"/api/v1/users/{new_id}", headers=admin, json={"is_active": False})
    assert deactivated.status_code == 200
    assert client.post("/api/v1/auth/login", json={"username": "new.viewer", "password": "New-viewer-password-123!"}).status_code == 403
    assert client.get("/api/v1/auth/me", headers=new_headers).status_code == 401

    with factory() as db:
        admin_id = db.scalar(select(User.id).where(User.username == "admin"))
    assert client.put(f"/api/v1/users/{admin_id}", headers=admin, json={"role": "USER"}).status_code == 409
    assert client.put(f"/api/v1/users/{admin_id}", headers=admin, json={"is_active": False}).status_code == 409
    assert client.put("/api/v1/users/9999", headers=admin, json={"full_name": "Ghost"}).status_code == 404

    reset = client.put(f"/api/v1/users/{new_id}", headers=admin, json={"is_active": True, "password": "Another-long-password-123!"})
    assert reset.status_code == 200
    assert client.post("/api/v1/auth/login", json={"username": "new.viewer", "password": "Another-long-password-123!"}).status_code == 200
    with factory() as db:
        actions = {row for row in db.scalars(select(AuditLog.action)).all()}
    assert {"USER_CREATE", "USER_UPDATE"} <= actions


def test_last_active_admin_cannot_be_demoted_by_another_admin(api):
    client, factory, login = api
    admin = _auth_headers(login("admin", "Admin-password-123!"))
    second = client.post("/api/v1/users", headers=admin, json={"username": "second.admin", "full_name": "Second", "role": "ADMIN", "password": "Second-admin-password-123!"}).json()["data"]
    second_headers = _auth_headers(client.post("/api/v1/auth/login", json={"username": "second.admin", "password": "Second-admin-password-123!"}).json()["data"]["access_token"])
    with factory() as db:
        admin_id = db.scalar(select(User.id).where(User.username == "admin"))
    assert client.put(f"/api/v1/users/{admin_id}", headers=second_headers, json={"is_active": False}).status_code == 200
    assert client.put(f"/api/v1/users/{second['id']}", headers=second_headers, json={"role": "USER"}).status_code == 409


def test_dashboard_rules_zero_tolerance_alokasi_pct_and_projection(api):
    client, factory, login = api
    token = login("viewer", "Viewer-password-123!")

    with factory.begin() as db:
        user_id = db.scalar(select(User.id).where(User.username == "admin"))
        department_id = db.scalar(select(Department.id).where(Department.code == "MIS000"))
        coas = {}
        for code, name, category in [
            ("100000001", "Exact", "Labor"), ("100000002", "Under", "Labor"), ("100000003", "Over", "Expenses"),
            ("100000004", "Allocation", "Expenses"), ("100000005", "Zero budget", "Expenses"),
        ]:
            coa = Coa(code=code, name=name, category=category, department_id=department_id, is_active=True, is_gl_derived=False)
            db.add(coa)
            db.flush()
            coas[code] = coa
        budget_batch = UploadBatch(kind="BUDGET", filename="b.xlsx", fy=2026, period=None, uploaded_by=user_id, rows_read=5, rows_imported=5, rows_rejected=0, total_amount=Decimal("0"), status="ACTIVE")
        gl_apr = UploadBatch(kind="GL", filename="apr.xlsx", fy=2026, period=1, uploaded_by=user_id, rows_read=1, rows_imported=1, rows_rejected=0, total_amount=Decimal("0"), status="ACTIVE")
        gl_jun = UploadBatch(kind="GL", filename="jun.xlsx", fy=2026, period=3, uploaded_by=user_id, rows_read=1, rows_imported=1, rows_rejected=0, total_amount=Decimal("0"), status="ACTIVE")
        replaced = UploadBatch(kind="GL", filename="old.xlsx", fy=2026, period=2, uploaded_by=user_id, rows_read=1, rows_imported=1, rows_rejected=0, total_amount=Decimal("0"), status="REPLACED")
        db.add_all([budget_batch, gl_apr, gl_jun, replaced])
        db.flush()

        def budget(code, per_month):
            for period in range(1, 13):
                db.add(BudgetEntry(fy=2026, period=period, coa_id=coas[code].id, amount=Decimal(per_month), batch_id=budget_batch.id))

        def actual(code, period, amount, batch):
            db.add(ActualEntry(fy=2026, period=period, coa_id=coas[code].id, amount=Decimal(amount), account_number=f"{code}-X-MIS000", section="MIS000", txn_date=None, description="t", currency="USD", exch_rate=Decimal("1"), debit_native=Decimal("0"), credit_native=Decimal("0"), row_no=2, batch_id=batch.id))

        budget("100000001", "10.00")
        budget("100000002", "10.00")
        budget("100000003", "0.50")
        budget("100000004", "-5.00")
        budget("100000005", "0.00")
        actual("100000001", 1, "10.004", gl_apr)  # rounds to exact match: 10.00
        actual("100000001", 3, "9.996", gl_jun)
        actual("100000002", 1, "5.00", gl_apr)
        actual("100000003", 1, "10.01", gl_apr)
        actual("100000004", 1, "1.00", gl_apr)
        actual("100000005", 3, "7.00", gl_jun)
        actual("100000003", 2, "999.00", replaced)  # replaced batch must be ignored

    response = client.get("/api/v1/dashboard?fiscal_year=2026", headers=_auth_headers(token))
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    accounts = {row["coa_code"]: row for row in data["accounts"]}

    assert accounts["100000001"]["budget"] == "120.00" and accounts["100000001"]["actual"] == "20.00"
    assert accounts["100000002"]["variance"] == "115.00" and accounts["100000002"]["status"] == "UNDER_BUDGET"
    assert accounts["100000003"]["actual"] == "10.01" and accounts["100000003"]["status"] == "OVER_BUDGET"
    assert accounts["100000004"]["status"] == "ALOKASI" and accounts["100000004"]["variance"] == "-61.00"
    assert accounts["100000005"]["variance_pct"] is None and accounts["100000005"]["status"] == "OVER_BUDGET"
    assert accounts["100000002"]["variance_pct"] == "95.83"
    months = accounts["100000003"]["monthly"]
    assert len(months) == 12
    assert months[0]["loaded"] is True and months[0]["actual"] == "10.01"
    assert months[1]["loaded"] is False and months[1]["actual"] == "0.00"
    assert months[2]["loaded"] is True and months[2]["actual"] == "0.00"
    assert accounts["100000004"]["monthly"][0]["budget"] == "-5.00"

    # Quarter filter: Q1 holds three budget months for the account (30.00) against 20.00 actual
    april = client.get("/api/v1/dashboard?fiscal_year=2026&quarter=Q1", headers=_auth_headers(token)).json()["data"]
    q1 = {row["coa_code"]: row for row in april["accounts"]}
    assert q1["100000001"]["status"] == "UNDER_BUDGET"  # Apr-Jun: 30.00 budget vs 20.00 actual

    summary = data["summary"]
    assert summary["alokasi_count"] == 1
    assert summary["over_budget_count"] == 2
    assert summary["under_budget_count"] == 2  # 100000001 (120 vs 20) and 100000002
    assert summary["on_budget_count"] == 0
    assert summary["budget"] == "186.00"  # 120 + 120 + 6 - 60 + 0
    assert summary["actual"] == "43.01"
    assert summary["variance"] == "142.99"

    projection = data["projection"]
    assert projection["months_loaded"] == 2 and projection["loaded_periods"] == [1, 3]
    assert projection["actual_to_date"] == "43.01"
    assert projection["year_end_actual"] == "258.06"  # 43.01 / 2 * 12, not / 3
    assert projection["remaining_budget"] == "142.99"
    assert projection["months_remaining"] == 10
    assert projection["allowed_monthly_spend"] == "14.30"

    labor = client.get("/api/v1/dashboard?fiscal_year=2026&category=Labor", headers=_auth_headers(token)).json()["data"]
    assert {row["coa_code"] for row in labor["accounts"]} == {"100000001", "100000002"}
    assert set(labor["available_categories"]) == {"Labor", "Expenses"}
    assert [row["quarter"] for row in data["quarters"]] == ["Q1", "Q2", "Q3", "Q4"]
    assert data["quarters"][0]["loaded"] is False  # Q1 has Apr and Jun but not May
    assert client.get("/api/v1/dashboard?fiscal_year=2026").status_code == 401
    assert client.get("/api/v1/dashboard?fiscal_year=2026&quarter=Q9", headers=_auth_headers(token)).status_code == 422


def _excel_budget_truth(root):
    """Independent reading of the dummy Budget workbook: FY'26 Budget (col G) summed per TOTAL block."""
    import re

    sheet = load_workbook(root / "Docs" / "Source" / "Budget Dummy.xlsx", read_only=True, data_only=True)["MIS (FC)"]
    categories: dict[str, Decimal] = {}
    block: list[tuple[str, Decimal]] = []
    seen: set[str] = set()
    per_code: dict[str, Decimal] = {}
    for row in sheet.iter_rows(min_row=7, values_only=True):
        description = str(row[1] or "").strip()
        if "total" in description.casefold():
            name = " ".join(w if w.upper() == "SGA" else w.title() for w in re.sub(r"^\s*total\s+", "", description, flags=re.I).strip().split())
            categories[name] = sum((amount for _code, amount in block), Decimal("0"))
            block = []
            continue
        code = str(row[0] or "").strip()
        if re.fullmatch(r"\d{9}", code) and code not in seen:
            seen.add(code)
            amount = Decimal(str(row[6] or 0))
            block.append((code, amount))
            per_code[code] = amount
    return categories, per_code


def test_dashboard_reconciles_with_dummy_excel_files(api):
    from collections import defaultdict
    from pathlib import Path

    client, factory, login = api
    auth = _auth_headers(login("admin", "Admin-password-123!"))
    viewer = _auth_headers(login("viewer", "Viewer-password-123!"))
    root = Path(__file__).resolve().parents[2]
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    budget_bytes = (root / "Docs" / "Source" / "Budget Dummy.xlsx").read_bytes()
    preview = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "BUDGET"}, files={"file": ("Budget Dummy.xlsx", budget_bytes, xlsx)}).json()
    assert preview["success"] is True, preview
    assert client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": preview["data"]["preview_id"], "decision": "CONFIRM"}).status_code == 200

    gl_bytes = (root / "Docs" / "Source" / "GL Dummy.xlsx").read_bytes()
    first = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("GL Dummy.xlsx", gl_bytes, xlsx)}).json()
    for error in first["errors"]:
        code = error["issue"].split()[2]
        assert client.post("/api/v1/coa", headers=auth, json={"code": code, "name": f"GL-derived {code}"}).status_code == 201
    gl = client.post("/api/v1/uploads/preview", headers=auth, data={"kind": "GL"}, files={"file": ("GL Dummy.xlsx", gl_bytes, xlsx)}).json()
    assert gl["success"] is True, gl
    assert client.post("/api/v1/uploads/confirm", headers=auth, json={"preview_id": gl["data"]["preview_id"], "decision": "CONFIRM"}).status_code == 200

    data = client.get("/api/v1/dashboard?fiscal_year=2026", headers=viewer).json()["data"]

    # Budget side must equal the Excel TOTAL rows block by block, and carry real categories.
    truth_categories, truth_codes = _excel_budget_truth(root)
    # Only the two accounts that exist in the GL but not in the budget may lack a category; the trailing allocation block is labelled.
    assert {row["coa_code"] for row in data["accounts"] if row["category"] == "Uncategorized"} == {"770102000", "770107001"}
    assert "SGA Labor" in {row["category"] for row in data["categories"]}
    allocation = next(row for row in data["categories"] if row["category"] == "Allocation")
    assert allocation["budget"] == "-368593.63" and allocation["status"] == "ALOKASI"
    shown = {row["category"]: Decimal(row["budget"]) for row in data["categories"]}
    for name, expected in truth_categories.items():
        if name in shown or expected != 0:
            assert shown[name] == expected.quantize(Decimal("0.01")), (name, shown.get(name), expected)
    assert Decimal(data["summary"]["budget"]) == sum(truth_codes.values(), Decimal("0")).quantize(Decimal("0.01"))

    # Actual side must equal the GL workbook: MIS000 rows, USD debit - credit, one period.
    # We include len==3 with MIS000, and len==2 if the code is known in the system.
    sheet = load_workbook(root / "Docs" / "Source" / "GL Dummy.xlsx", read_only=True, data_only=True)["CORE"]
    gl_total = Decimal("0")
    gl_codes: dict[str, Decimal] = defaultdict(Decimal)
    by_code = {row["coa_code"]: row for row in data["accounts"]}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        account = str(row[3] or "").strip()
        parts = account.split("-")
        if not parts:
            continue
        code = parts[0].strip()
        in_scope = len(parts) == 3 and parts[2].strip().upper() == "MIS000"
        section_missing = len(parts) == 2 and code in by_code
        if in_scope or section_missing:
            value = Decimal(str(row[13] or 0)) - Decimal(str(row[14] or 0))
            gl_total += value
            gl_codes[code] += value
    assert Decimal(data["summary"]["actual"]) == gl_total.quantize(Decimal("0.01"))
    for code, value in gl_codes.items():
        assert Decimal(by_code[code]["actual"]) == value.quantize(Decimal("0.01")), code

    # Variance is budget - actual everywhere, and the month the GL belongs to is the only loaded one.
    assert Decimal(data["summary"]["variance"]) == Decimal(data["summary"]["budget"]) - Decimal(data["summary"]["actual"])
    assert data["projection"]["months_loaded"] == 1 and data["projection"]["loaded_periods"] == [3]
    assert data["monthly"][2]["actual"] == data["summary"]["actual"]
    assert sum(Decimal(row["actual"]) for row in data["monthly"]) == Decimal(data["summary"]["actual"])
    assert data["accounts"] and all(row["category"] for row in data["accounts"])

def test_status_is_zero_tolerance_after_single_rounding():
    from backend.vega.routers.dashboard import _pct, _status

    assert _status(Decimal("10.00"), Decimal("10.004")) == "ON_BUDGET"
    assert _status(Decimal("10.00"), Decimal("10.006")) == "OVER_BUDGET"
    assert _status(Decimal("10.00"), Decimal("9.99")) == "UNDER_BUDGET"
    assert _status(Decimal("-0.01"), Decimal("100")) == "ALOKASI"
    assert _status(Decimal("0"), Decimal("0")) == "ON_BUDGET"
    assert _pct(Decimal("5"), Decimal("0")) is None
    assert _pct(Decimal("-5"), Decimal("-10")) == "-50.00"


def test_logout_revokes_only_current_session_and_survives_app_recreation(api):
    client, factory, login = api
    first = _auth_headers(login("viewer", "Viewer-password-123!"))
    second = _auth_headers(login("viewer", "Viewer-password-123!"))
    assert first != second
    assert client.post("/api/v1/auth/logout", headers=first).status_code == 200
    assert client.get("/api/v1/auth/me", headers=first).status_code == 401
    assert client.get("/api/v1/dashboard?fiscal_year=2026", headers=first).status_code == 401
    assert client.get("/api/v1/auth/me", headers=second).status_code == 200
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "LOGOUT")) == 1
    another = create_app()
    another.dependency_overrides.update(client.app.dependency_overrides)
    with TestClient(another) as restarted:
        assert restarted.get("/api/v1/auth/me", headers=first).status_code == 401
        assert restarted.get("/api/v1/auth/me", headers=second).status_code == 200


def test_password_change_reset_and_deactivation_revoke_sessions(api):
    client, factory, login = api
    admin = _auth_headers(login("admin", "Admin-password-123!"))
    viewer = _auth_headers(login("viewer", "Viewer-password-123!"))
    assert client.put("/api/v1/auth/me/password", headers=viewer, json={"old_password":"wrong", "new_password":"Changed-password-123!"}).status_code == 422
    assert client.put("/api/v1/auth/me/password", headers=viewer, json={"old_password":"Viewer-password-123!", "new_password":"Changed-password-123!"}).status_code == 200
    assert client.get("/api/v1/auth/me", headers=viewer).status_code == 401
    assert client.post("/api/v1/auth/login", json={"username":"viewer","password":"Viewer-password-123!"}).status_code == 401
    fresh = _auth_headers(login("viewer", "Changed-password-123!"))
    viewer_id = client.get("/api/v1/auth/me", headers=fresh).json()["data"]["id"]
    assert client.put(f"/api/v1/users/{viewer_id}", headers=admin, json={"password":"Reset-password-123!"}).status_code == 200
    assert client.get("/api/v1/auth/me", headers=fresh).status_code == 401
    new = _auth_headers(login("viewer", "Reset-password-123!"))
    assert client.put(f"/api/v1/users/{viewer_id}", headers=admin, json={"is_active":False}).status_code == 200
    assert client.get("/api/v1/auth/me", headers=new).status_code == 401
    assert client.put(f"/api/v1/users/{viewer_id}", headers=admin, json={"is_active":True}).status_code == 200
    assert client.get("/api/v1/auth/me", headers=new).status_code == 401
    assert client.get("/api/v1/auth/me", headers=_auth_headers(login("viewer", "Reset-password-123!"))).status_code == 200


def test_invalid_expired_tampered_tokens_and_current_db_role(api):
    import jwt
    from backend.vega.config import jwt_secret
    from backend.vega.security import decode_access_token
    client, factory, login = api
    token = login("admin", "Admin-password-123!")
    payload = decode_access_token(token)
    payload["exp"] = 1
    expired = jwt.encode(payload, jwt_secret(), algorithm="HS256")
    for invalid in ("invalid", token + "tamper", expired):
        assert client.get("/api/v1/auth/me", headers=_auth_headers(invalid)).status_code == 401
    with factory.begin() as db:
        db.scalar(select(User).where(User.username == "admin")).role = "USER"
    assert client.get("/api/v1/users", headers=_auth_headers(token)).status_code == 403
    assert client.get("/api/v1/auth/me", headers=_auth_headers(token)).json()["data"]["role"] == "USER"


def test_all_write_endpoints_deny_viewer_and_anonymous(api):
    client, _, login = api
    viewer = _auth_headers(login("viewer", "Viewer-password-123!"))
    writes = [
        ("POST", "/api/v1/coa", {"code":"999888777","name":"Denied"}),
        ("PUT", "/api/v1/coa/1", {"name":"Denied"}),
        ("PATCH", "/api/v1/coa/1/status", {"is_active":False}),
        ("DELETE", "/api/v1/coa/1", None),
        ("POST", "/api/v1/users", {"username":"denied","full_name":"Denied","role":"ADMIN","password":"Denied-password-123!"}),
        ("PUT", "/api/v1/users/1", {"role":"ADMIN"}),
        ("POST", "/api/v1/uploads/confirm", {"preview_id":"unknown-preview-1234567890","decision":"CONFIRM"}),
        ("POST", "/api/v1/uploads/batches/1/cancel", None),
    ]
    for method, path, body in writes:
        assert client.request(method, path, headers=viewer, json=body).status_code == 403, path
        assert client.request(method, path, json=body).status_code == 401, path


def test_password_byte_limit_and_coa_pagination(api):
    client, factory, login = api
    auth = _auth_headers(login("admin", "Admin-password-123!"))
    for password in ("Aa1!" + "x"*70, "Aa1!" + "é"*35):
        result = client.post("/api/v1/users",headers=auth,json={"username":"long-password","full_name":"Long","role":"USER","password":password})
        assert result.status_code == 422
    with factory.begin() as db:
        dept = db.scalar(select(Department.id))
        db.add_all([Coa(code=str(900000000+i), name=f"Pagination {i}",department_id=dept,is_active=True,is_gl_derived=False) for i in range(105)])
    one=client.get("/api/v1/coa?page=1&page_size=100",headers=auth).json()["data"]
    two=client.get("/api/v1/coa?page=2&page_size=100",headers=auth).json()["data"]
    assert one["total"] == two["total"] == 105
    assert len(one["items"]) == 100 and len(two["items"]) == 5
    assert not ({a["id"] for a in one["items"]} & {a["id"] for a in two["items"]})
    assert client.get("/api/v1/coa?page=0",headers=auth).status_code == 422
    assert client.get("/api/v1/uploads/batches?page_size=101",headers=auth).status_code == 422
    assert client.get("/api/v1/coa/99999",headers=auth).status_code == 404
    assert client.get("/api/v1/uploads/batches/99999",headers=auth).status_code == 404


def test_preview_owner_expiry_replay_and_replacement_backup_failure(api, tmp_path, monkeypatch):
    from backend.vega.backup import BackupError
    client, factory, login = api
    admin = _auth_headers(login("admin", "Admin-password-123!"))
    assert client.post("/api/v1/coa",headers=admin,json={"code":"999999999","name":"QA"}).status_code == 201
    path = _make_gl_file(tmp_path / "qa.xlsx")
    def preview():
        return client.post("/api/v1/uploads/preview",headers=admin,data={"kind":"GL"},files={"file":(path.name,path.read_bytes())}).json()["data"]["preview_id"]
    owner = client.post("/api/v1/users",headers=admin,json={"username":"other.admin","full_name":"Other","role":"ADMIN","password":"Other-password-123!"}).json()["data"]
    other = _auth_headers(login("other.admin","Other-password-123!"))
    pid=preview()
    assert client.post("/api/v1/uploads/confirm",headers=other,json={"preview_id":pid,"decision":"CONFIRM"}).status_code == 410
    uploads._previews[pid]["created_at"] -= uploads.PREVIEW_TTL_SECONDS + 1
    assert client.post("/api/v1/uploads/confirm",headers=admin,json={"preview_id":pid,"decision":"CONFIRM"}).status_code == 410
    pid=preview()
    assert client.post("/api/v1/uploads/confirm",headers=admin,json={"preview_id":pid,"decision":"CONFIRM"}).status_code == 200
    assert client.post("/api/v1/uploads/confirm",headers=admin,json={"preview_id":pid,"decision":"CONFIRM"}).status_code == 410
    pid=preview()
    def fail_backup(*args,**kwargs): raise BackupError("simulated")
    monkeypatch.setattr(uploads,"create_database_backup",fail_backup)
    assert client.post("/api/v1/uploads/confirm",headers=admin,json={"preview_id":pid,"decision":"REPLACE"}).status_code == 503
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(ActualEntry)) == 1
        assert db.scalar(select(func.sum(ActualEntry.amount))) == Decimal("12.5")
        assert db.scalar(select(func.count()).select_from(UploadBatch).where(UploadBatch.status == "ACTIVE")) == 1


def test_whitespace_names_and_invalid_file_inputs(api):
    client, _, login = api
    h=_auth_headers(login("admin","Admin-password-123!"))
    assert client.post("/api/v1/users",headers=h,json={"username":"blank","full_name":"   ","role":"USER","password":"Valid-password-123!"}).status_code == 422
    assert client.post("/api/v1/coa",headers=h,json={"code":"999111222","name":"   "}).status_code == 422
    for name,blob in [("empty.xlsx",b""),("legacy.xls",b"bad"),("corrupt.xlsx",b"bad")]:
        r=client.post("/api/v1/uploads/preview",headers=h,data={"kind":"GL"},files={"file":(name,blob)})
        assert r.status_code == 200 and r.json()["success"] is False


def test_dashboard_uncategorized_filter_and_net_zero_monthly_account(api):
    client, factory, login = api
    h=_auth_headers(login("viewer","Viewer-password-123!"))
    with factory.begin() as db:
        uid=db.scalar(select(User.id).where(User.username == "admin")); dept=db.scalar(select(Department.id))
        coa=Coa(code="999111222",name="Offset",department_id=dept,is_active=True,is_gl_derived=False); db.add(coa); db.flush()
        batch=UploadBatch(kind="BUDGET",filename="offset.xlsx",fy=2026,period=None,uploaded_by=uid,rows_read=1,rows_imported=1,rows_rejected=0,total_amount=Decimal("0"),status="ACTIVE"); db.add(batch); db.flush()
        db.add_all([BudgetEntry(fy=2026,period=1,coa_id=coa.id,amount=Decimal("10"),batch_id=batch.id),BudgetEntry(fy=2026,period=2,coa_id=coa.id,amount=Decimal("-10"),batch_id=batch.id)])
    data=client.get("/api/v1/dashboard?fiscal_year=2026&category=Uncategorized",headers=h).json()["data"]
    assert "Uncategorized" in data["available_categories"]
    assert data["accounts"][0]["has_data"] is True
    assert data["accounts"][0]["budget"] == "0.00"


def test_backup_decodes_url_credentials(tmp_path, monkeypatch):
    import subprocess
    from pathlib import Path
    from backend.vega import backup
    seen={}
    monkeypatch.setattr(backup.shutil,"which",lambda _: "pg_dump")
    def fake(cmd,**kwargs):
        seen.update(kwargs["env"])
        Path(cmd[cmd.index("--file")+1]).write_text("-- PostgreSQL database dump\n",encoding="utf-8")
        return subprocess.CompletedProcess(cmd,0,stdout="",stderr="")
    monkeypatch.setattr(backup.subprocess,"run",fake)
    backup.create_database_backup("postgresql://qa:p%40ss%3Aword@localhost/test",str(tmp_path))
    assert seen["PGPASSWORD"] == "p@ss:word"
