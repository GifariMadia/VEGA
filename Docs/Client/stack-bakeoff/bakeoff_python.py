"""Opsi A/C — read the GL with openpyxl.

Identical task to bakeoff_node.js: read the budget master, read GL sheet CORE,
apply the two-clause row filter, sum converted debit minus credit.

    python bakeoff_python.py <source-dir> [--detail] [--out FILE.csv|FILE.json]
"""
import json
import sys
import time
import tracemalloc
from openpyxl import load_workbook

SRC = sys.argv[1]
DETAIL = "--detail" in sys.argv
OUT = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else None

t0 = time.perf_counter()
tracemalloc.start()

# --- 1. budget master -> set of COA codes (Excel-Template-Spec 2.3, 2.4) ---
wb = load_workbook(f"{SRC}/Budget Dummy.xlsx", read_only=True, data_only=True)
budget_coa = set()
for row in wb["MIS (FC)"].iter_rows(min_row=7, values_only=True):
    desc = row[1]
    if not desc or "TOTAL" in str(desc).upper():
        continue                                   # subtotal row
    code = str(row[0]).strip()                     # column A mixes str and int
    if code.isdigit() and len(code) == 9:
        budget_coa.add(code)
wb.close()

# --- 2. GL sheet CORE ---
wb = load_workbook(f"{SRC}/GL Dummy.xlsx", read_only=True, data_only=True)
gl = wb["CORE"].iter_rows(values_only=True)

header = next(gl)
if header[11:15] != ("Debits", "Credits", "Debits", "Credits"):
    sys.exit(f"REJECT: unexpected header in columns L-O: {header[11:15]}")

rows, skipped = [], 0
for r in gl:
    acct, period = r[3], r[0]
    if not acct or not period:                     # rules 1 & 2 -> the junk row dies here
        skipped += 1
        continue
    parts = str(acct).split("-")
    if len(parts) == 3 and parts[2] == "MIS000":
        clause = "A"
    elif len(parts) == 2 and parts[0] in budget_coa:
        clause = "B"
    else:
        continue
    debit, credit = r[13] or 0, r[14] or 0         # by index, never by header name
    if not isinstance(debit, (int, float)) or not isinstance(credit, (int, float)):
        sys.exit(f"REJECT: non-numeric amount at {acct}")           # rule 3
    rows.append({
        "account": str(acct), "coa": parts[0], "description": str(r[4] or ""),
        "curr": str(r[9] or ""), "debit": debit, "credit": credit,
        "net": debit - credit, "clause": clause,          # raw; rounded only on output
    })
wb.close()

elapsed = time.perf_counter() - t0
peak_mb = tracemalloc.get_traced_memory()[1] / 1e6
total = sum(x["net"] for x in rows)                       # round once, at the end
by_clause = {c: sum(1 for x in rows if x["clause"] == c) for c in "AB"}

# --- 3. output ---
print(f"python/openpyxl | clause_a={by_clause['A']} clause_b={by_clause['B']} "
      f"total_rows={len(rows)} | sum_usd={total:.2f} | "
      f"skipped={skipped} | {elapsed:.2f}s | peak heap {peak_mb:.0f} MB")

if DETAIL:
    for clause, title in (("A", "3 segments, section MIS000"), ("B", "2 segments, COA in budget master")):
        hits = [x for x in rows if x["clause"] == clause]
        print(f"\nCLAUSE {clause} - {title}  ({len(hits)} rows)")
        for x in hits:
            print(f"  {x['account']:<26} {x['description'][:34]:<34} {x['curr']:<4} {x['net']:>12,.2f}")
    print(f"\nSUBTOTAL PER COA")
    subtotal = {}
    for x in rows:
        subtotal[x["coa"]] = subtotal.get(x["coa"], 0) + x["net"]
    for coa in sorted(subtotal):
        print(f"  {coa}  {subtotal[coa]:>14,.2f}")
    print(f"  {'TOTAL':<9}  {total:>14,.2f}")
    print(f"\nSKIPPED: {skipped} row(s) - no account number and/or no period")

if OUT:
    # money is formatted to 2 dp here, on the way out, and nowhere earlier
    out_rows = [{**x, "debit": f"{x['debit']:.2f}", "credit": f"{x['credit']:.2f}",
                 "net": f"{x['net']:.2f}"} for x in rows]
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        if OUT.endswith(".json"):
            json.dump({"engine": "python/openpyxl", "seconds": round(elapsed, 2),
                       "peak_heap_mb": round(peak_mb), "packages": 1, "helper_lines": 0,
                       "total": f"{total:.2f}", "skipped": skipped, "rows": out_rows}, f, indent=1)
        else:
            f.write("account,description,curr,debit,credit,net,clause\n")
            for x in out_rows:
                desc = x["description"].replace('"', '""')
                f.write(f'{x["account"]},"{desc}",{x["curr"]},{x["debit"]},{x["credit"]},{x["net"]},{x["clause"]}\n')
    print(f"wrote {OUT}")
